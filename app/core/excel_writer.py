"""Excel writer for fixed SPED layouts."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TEXT_COLUMNS = {
    "CNPJ",
    "Código Participante",
    "CNPJ Participante",
    "CPF Participante",
    "CNPJ/CPF Participante",
    "Número Documento",
    "Série",
    "Modelo",
    "Chave NF-e",
    "Código Item",
    "Código Barra",
    "NCM",
    "CFOP",
    "CST ICMS",
    "CST IPI",
    "CST PIS",
    "CST Cofins",
    "Conta Contábil",
}

DATE_COLUMNS = {"Período", "Data Documento", "Data Entrada/Saída"}


def aplicar_formatacao_excel(ws, cabecalho: list[str], total_linhas: int) -> None:
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="D9EAF7")

    for coluna, nome in enumerate(cabecalho, start=1):
        cell = ws.cell(row=1, column=coluna)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        cell.fill = fill
        letra = get_column_letter(coluna)
        largura = min(max(len(nome) + 2, 12), 42)
        ws.column_dimensions[letra].width = largura

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecalho))}{max(total_linhas + 1, 1)}"

    indice_por_nome = {nome: idx for idx, nome in enumerate(cabecalho, start=1)}
    for nome in DATE_COLUMNS & set(cabecalho):
        letra = get_column_letter(indice_por_nome[nome])
        for cell in ws[letra][1:]:
            cell.number_format = "dd/mm/yyyy"


def salvar_excel(caminho_saida: str | Path, cabecalho: list[str], linhas: Iterable[dict]) -> Path:
    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha"
    ws.append(cabecalho)

    total_linhas = 0
    for linha in linhas:
        valores = []
        for coluna in cabecalho:
            valor = linha.get(coluna, "")
            if isinstance(valor, Decimal):
                valor = float(valor)
            valores.append(valor)
        ws.append(valores)
        total_linhas += 1

    aplicar_formatacao_excel(ws, cabecalho, total_linhas)
    wb.save(caminho_saida)
    wb.close()
    return caminho_saida
